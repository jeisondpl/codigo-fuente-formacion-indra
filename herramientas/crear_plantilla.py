# -*- coding: utf-8 -*-
"""Genera la plantilla Excel de usuarios para el Portal Formación Indra."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

wb = Workbook()

AZUL = "0A3D62"
GRIS = "F2F4F7"
BORDE = Border(*[Side(style="thin", color="D0D5DD")] * 4)

# ---------------------------------------------------------------- Listas ----
ws_listas = wb.create_sheet("Listas")
ws_listas["A1"] = "ROL"
ws_listas["B1"] = "CAMPAÑA"
ws_listas["C1"] = "SUPERVISORES (editar esta lista libremente)"
for c in ("A1", "B1", "C1"):
    ws_listas[c].font = Font(bold=True, color="FFFFFF")
    ws_listas[c].fill = PatternFill("solid", fgColor=AZUL)

roles = ["Asesor", "Supervisor", "Formador"]
campanas = ["Clientes Críticos", "STA", "Ambas"]
for i, v in enumerate(roles, start=2):
    ws_listas[f"A{i}"] = v
for i, v in enumerate(campanas, start=2):
    ws_listas[f"B{i}"] = v
# Supervisores conocidos (el formador edita/agrega aquí)
supervisores = ["(agregue aquí los supervisores)"]
for i, v in enumerate(supervisores, start=2):
    ws_listas[f"C{i}"] = v

ws_listas.column_dimensions["A"].width = 14
ws_listas.column_dimensions["B"].width = 18
ws_listas.column_dimensions["C"].width = 46
RANGO_ROL = "Listas!$A$2:$A$200"
RANGO_CAMPANA = "Listas!$B$2:$B$200"
RANGO_SUPERVISOR = "Listas!$C$2:$C$200"

# -------------------------------------------------------------- Usuarios ----
ws = wb.active
ws.title = "Usuarios"
headers = ["DOCUMENTO", "NOMBRE", "ROL", "CAMPAÑA", "SUPERVISOR"]
for col, h in enumerate(headers, start=1):
    c = ws.cell(row=1, column=col, value=h)
    c.font = Font(bold=True, color="FFFFFF", size=11)
    c.fill = PatternFill("solid", fgColor=AZUL)
    c.alignment = Alignment(horizontal="center")

# Ejemplos
ejemplos = [
    ["1030123456", "Pérez Gómez, Ana María", "Asesor", "Clientes Críticos", ""],
    ["1030654321", "Ramírez Ruiz, Luis", "Asesor", "Ambas", ""],
    ["1030987654", "García López, Carolina", "Supervisor", "Clientes Críticos", ""],
]
for r, fila in enumerate(ejemplos, start=2):
    for col, v in enumerate(fila, start=1):
        ws.cell(row=r, column=col, value=v)

anchos = [16, 34, 14, 20, 30]
for i, a in enumerate(anchos, start=1):
    ws.column_dimensions[get_column_letter(i)].width = a
for row in ws.iter_rows(min_row=2, max_row=500, max_col=5):
    for cell in row:
        cell.border = BORDE
        if cell.column == 1:
            cell.number_format = "@"  # documento como texto

# Desplegables sobre las columnas C (ROL), D (CAMPAÑA), E (SUPERVISOR)
dv_rol = DataValidation(type="list", formula1=RANGO_ROL, allow_blank=True,
                        showErrorMessage=True, errorTitle="Valor no válido",
                        error="Seleccione: Asesor, Supervisor o Formador")
dv_cam = DataValidation(type="list", formula1=RANGO_CAMPANA, allow_blank=True,
                        showErrorMessage=True, errorTitle="Valor no válido",
                        error="Seleccione: Clientes Críticos, STA o Ambas")
dv_sup = DataValidation(type="list", formula1=RANGO_SUPERVISOR, allow_blank=True,
                        showErrorMessage=False)  # permite nombres nuevos fuera de lista
ws.add_data_validation(dv_rol)
ws.add_data_validation(dv_cam)
ws.add_data_validation(dv_sup)
dv_rol.add("C2:C500")
dv_cam.add("D2:D500")
dv_sup.add("E2:E500")
ws.freeze_panes = "A2"

# --------------------------------------------------------- Instrucciones ----
wi = wb.create_sheet("Instrucciones")
lineas = [
    ("CÓMO USAR ESTA PLANTILLA — Portal Formación Indra", True),
    ("", False),
    ("1. Diligencie una fila por usuario en la hoja «Usuarios».", False),
    ("2. DOCUMENTO: cédula sin puntos ni comas (queda como texto).", False),
    ("3. NOMBRE: nombre completo del usuario.", False),
    ("4. ROL: desplegable (Asesor / Supervisor / Formador).", False),
    ("5. CAMPAÑA: desplegable (Clientes Críticos / STA / Ambas).", False),
    ("   - Si el rol es Formador, la campaña se ignora: siempre ve las dos.", False),
    ("6. SUPERVISOR: desplegable con los nombres de la hoja «Listas».", False),
    ("   - Solo aplica para asesores. Deje vacío si no tiene supervisor.", False),
    ("   - El nombre debe coincidir con un usuario Supervisor ya creado.", False),
    ("7. En la app (Panel del Formador) use «Importar Excel» para crearlos", False),
    ("   todos de una vez; los que ya existan se actualizan, no se duplican.", False),
    ("8. También puede usar «Descargar Excel» en la app para obtener el", False),
    ("   listado actual de la base de datos en este mismo formato.", False),
    ("", False),
    (" Las filas de ejemplo se pueden borrar antes de importar.", True),
]
for r, (texto, negrita) in enumerate(lineas, start=1):
    c = wi.cell(row=r, column=1, value=texto)
    c.font = Font(bold=negrita, size=12 if negrita else 11)
wi.column_dimensions["A"].width = 70

wb.save("Plantilla_Usuarios_Formacion_Indra.xlsx")
print("Plantilla creada")
